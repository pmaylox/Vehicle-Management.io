from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def style_sheet(sheet, title, headers, rows, widths, formats):
    column_count = len(headers)

    sheet.sheet_view.showGridLines = False
    sheet.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=column_count
    )

    sheet.cell(1, 1, title)
    sheet.cell(1, 1).font = Font(
        name="Calibri", size=18, bold=True, color="FFFFFF"
    )
    sheet.cell(1, 1).fill = PatternFill(
        "solid", fgColor="17365D"
    )
    sheet.row_dimensions[1].height = 36

    sheet.cell(2, 1, f"Exported {date.today():%m/%d/%Y}")
    sheet.cell(2, 1).font = Font(
        name="Calibri", size=10, color="666666"
    )

    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(4, column, header)
        cell.font = Font(
            name="Calibri", size=12, bold=True, color="FFFFFF"
        )
        cell.fill = PatternFill("solid", fgColor="245A81")
        cell.alignment = Alignment(
            horizontal="center",
            vertical="top", wrap_text=True
        )

    sheet.row_dimensions[4].height = 36

    for row_number, values in enumerate(rows, start=5):
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row_number, column, value)

            # Keep descriptions as text, even if they start with "=".
            if isinstance(value, str):
                cell.data_type = "s"

            cell.font = Font(name="Calibri", size=11)
            cell.alignment = Alignment(
                vertical="top", wrap_text=True
            )

            if row_number % 2 == 1:
                cell.fill = PatternFill(
                    "solid", fgColor="EDF3F8"
                )

            if column in formats:
                cell.number_format = formats[column]

        sheet.row_dimensions[row_number].height = 45

    for column, width in enumerate(widths, start=1):
        letter = get_column_letter(column)
        sheet.column_dimensions[letter].width = width

    last_column = get_column_letter(column_count)
    last_row = max(4, sheet.max_row)

    sheet.freeze_panes = "A5"
    sheet.auto_filter.ref = f"A4:{last_column}{last_row}"
    sheet.print_title_rows = "1:4"
    sheet.oddFooter.center.text = "Page &P of &N"
    sheet.oddFooter.center.size = 10
    sheet.oddFooter.left.text = title
    sheet.oddFooter.left.size = 9
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0

    if not rows:
        sheet.cell(5, 1, "No records yet.")

def style_summary(sheet, start_row, row_count, row_height=32):
    for row in range(start_row, start_row + row_count):
        sheet.row_dimensions[row].height = row_height

        for column in (3, 4):
            cell = sheet.cell(row, column)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="EDF3F8")
            cell.alignment = Alignment(
                vertical="center", wrap_text=True
            )
   
    sheet.print_area = f"A1:{get_column_letter(sheet.max_column)}{start_row + row_count - 1}"

def export_vehicle_excel(vehicle, output_file=None):
    if output_file is None:
        safe_name = "".join(
            character if character.isalnum() else "_"
            for character in vehicle.name
        ).strip("_") or "vehicle"

        folder = Path(__file__).parent / "exports"
        output_file = folder / f"{vehicle.year}_{safe_name}.xlsx"

        if output_file.exists():
            confirm = input(
                f"{output_file.name} already exists. "
                "Replace it? (yes/no): "
            ).strip().lower()

            if confirm != "yes":
                print("Export canceled.")
                return

    output_file = Path(output_file)
    folder = output_file.parent

    workbook = Workbook()
    service_sheet = workbook.active
    service_sheet.title = "Service History"
    fuel_sheet = workbook.create_sheet("Fuel History")

    service_rows = []
    total_service_cost = Decimal("0")
    missing_service_cost = 0

    for record in vehicle.service_history:
        cents = record.get("cost_cents")
        if cents is not None:
            total_service_cost += Decimal(cents) / 100
        else:
            missing_service_cost += 1

        cost = (
            Decimal(cents) / 100
            if cents is not None
            else "Not recorded"
        )

        service_rows.append([
            date.fromisoformat(record["date"]),
            record["mileage"],
            record["description"],
            cost,
            (
                "Yes" if record["oil_reset"] else "No"
            ) if "oil_reset" in record else "Unknown",
            (
                "Yes" if record["tire_reset"] else "No"
            ) if "tire_reset" in record else "Unknown"
        ])

    total_fuel_cost = Decimal("0")
    total_gallons = Decimal("0")

     
    fuel_rows = []
    previous_full_mileage = None
    gallons_since_full = Decimal("0")

    for record in vehicle.fuel_history:
        gallons = Decimal(record["gallons"])
        total_gallons += gallons
        cost = Decimal(record["cost_cents"]) / 100
        total_fuel_cost += cost
        mpg = None

        if previous_full_mileage is not None:
            gallons_since_full += gallons

        if record["full_tank"]:
            if previous_full_mileage is not None:
                miles_driven = record["mileage"] - previous_full_mileage
                mpg = Decimal(miles_driven) / gallons_since_full

            previous_full_mileage = record["mileage"]
            gallons_since_full = Decimal("0")

        fuel_rows.append([
            date.fromisoformat(record["date"]),
            record["mileage"],
            gallons,
            cost,
            "Yes" if record["full_tank"] else "No",
            mpg,
            cost / gallons
        ])

    title = f"{vehicle.year} {vehicle.name}"

    style_sheet(
        service_sheet,
        f"{title} — Service History",
        [
            "Service Date", "Odometer (mi)", "Service Performed",
            "Cost", "Oil Reset", "Tire Reset"
        ],
        service_rows,
        [16, 19, 55, 18, 16, 16],
        {1: "mm/dd/yyyy", 2: "#,##0", 4: '"$"#,##0.00'}
    )

    style_sheet(
        fuel_sheet,
        f"{title} — Fuel History",
                [
            "Fill-up Date", "Odometer (mi)", "US Gallons",
            "Total Cost", "Full Tank", "-MPG", "Price/Gallon"
        ],
                fuel_rows,
        [18, 21, 19, 20, 18, 14, 18],
        {
            1: "mm/dd/yyyy",
            2: "#,##0",
            3: "0.000",
            4: '"$"#,##0.00',
            6: "0.00",
            7: '"$"0.000'
        }
    )
    
    service_with_cost = len(service_rows) - missing_service_cost
    average_service_cost = total_service_cost / service_with_cost if service_with_cost > 0 else None


    service_summary_row = service_sheet.max_row + 2
    service_sheet.cell(service_summary_row, 3, "Total recorded service spending")
    service_sheet.cell(service_summary_row, 4, total_service_cost)
    service_sheet.cell(service_summary_row, 4).number_format = '"$"#,##0.00'    
    service_sheet.cell(service_summary_row +1, 3, "Services with no cost recorded")
    service_sheet.cell(service_summary_row + 1, 4,  missing_service_cost)    
    service_sheet.cell(service_summary_row +2, 3, "Service records")
    service_sheet.cell(service_summary_row +2, 4, len(service_rows))   
    service_sheet.cell(service_summary_row +3, 3, "Average recorded service cost")
    service_sheet.cell(service_summary_row +3, 4, average_service_cost)   
    service_sheet.cell(service_summary_row +3, 4).number_format = '"$"#,##0.00'
    
    style_summary(service_sheet, service_summary_row, 4)
            
    summary_row = fuel_sheet.max_row + 2
    fuel_sheet.cell(summary_row, 3, "Total fuel spending")
    fuel_sheet.cell(summary_row, 4, total_fuel_cost)
    fuel_sheet.cell(summary_row, 4).number_format = '"$"#,##0.00'
    fuel_sheet.cell(summary_row +1, 3, "Total gallons")
    fuel_sheet.cell(summary_row +1, 4, total_gallons)
    fuel_sheet.cell(summary_row +1, 4).number_format = "0.000"
    average_price = total_fuel_cost / total_gallons if total_gallons > 0 else None
    fuel_sheet.cell(summary_row +2, 3, "Average price/gallon")
    fuel_sheet.cell(summary_row +2, 4, average_price)
    fuel_sheet.cell(summary_row +2, 4).number_format = '"$"0.000'
    fuel_sheet.cell(summary_row + 3, 3, "Fill-ups")
    fuel_sheet.cell(summary_row + 3, 4, len(fuel_rows))

    style_summary(fuel_sheet, summary_row, 4, row_height=40)

    try:
        folder.mkdir(exist_ok=True)
        workbook.save(output_file)
        print(f"Excel workbook exported to:\n{output_file}")
        return output_file

    except OSError as error:
        print(f"Could not export workbook: {error}")
        print("Close the workbook in Excel and try again.")

        return None

    finally:
        workbook.close()